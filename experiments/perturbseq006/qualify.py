"""Source-bound data and entrypoint qualification; never a scientific success judge."""
import argparse
import hashlib
import importlib.metadata
import json
from pathlib import Path
import resource
import subprocess
import sys
import time
from range_hdf5 import RangeFile

URL = 'https://huggingface.co/datasets/arcinstitute/Replogle-Nadig-Preprint/resolve/63a2f314b4043b002965d13cef84b4581b585601/rpe1_normalized_singlecell_01.h5ad'
PSGRN = '74aa640f7c472b23a69811f6795bb17678efd344'
LINEAR = 'bfa6eeea2bd145a1af2ec0127a2e808cc38456a9'
XLS_SHA = '1b652c352d4dabdd20324d7fd74d331898fa22c51cf6fa28ff65aa1322a6c78f'

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def pinned_source(root, commit):
    actual = subprocess.check_output(['git', '-C', str(root), 'rev-parse', 'HEAD'], text=True).strip()
    if actual != commit or subprocess.check_output(['git', '-C', str(root), 'status', '--porcelain']).strip():
        raise ValueError('UPSTREAM_SOURCE_NOT_EXACT')
    return actual

def run(output, psgrn_source, linear_source):
    import h5py
    import numpy as np
    import openpyxl
    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    started, cpu = time.monotonic(), time.process_time()
    report = {'schema': 'agi-alpha-data-qualification006/v1', 'scientific_verdict': 'NO_SCIENTIFIC_ADVANTAGE_ESTABLISHED',
              'paid_model_calls': 0, 'complete_cost': None, 'hypothesis': 'Source-bound range reads can remove whole-file storage for inspection; this is ordinary library composition, not new science.'}
    try:
        pinned_source(psgrn_source, PSGRN)
        pinned_source(linear_source, LINEAR)
        report['upstream'] = {'psgrn_commit': PSGRN, 'linear_paper_commit': LINEAR,
                              'main_exists': (psgrn_source/'src/main.py').is_file(),
                              'main_fast_exists': (psgrn_source/'src/main_fast.py').is_file()}
        p = subprocess.run([sys.executable, str(psgrn_source/'causalscbench/apps/main_app.py'), '--help'],
                           cwd=psgrn_source, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=30,
                           env={'PATH': '/usr/bin:/bin', 'PYTHONPATH': str(psgrn_source), 'LANG': 'C.UTF-8'})
        (output/'entrypoint.log').write_bytes(p.stdout)
        report['entrypoint'] = {'command': [sys.executable, 'causalscbench/apps/main_app.py', '--help'], 'returncode': p.returncode,
                                 'log_sha256': sha(output/'entrypoint.log'), 'qualification': 'UNQUALIFIED_MISSING_DEPENDENCIES' if p.returncode else 'IMPORT_INTERFACE_ONLY'}
        with RangeFile(URL, 8700873216) as remote:
            with h5py.File(remote, 'r') as f:
                cats=f['obs/__categories/gene_id'][:]
                codes=f['obs/gene_id'][:]
                varids=set(f['var/gene_id'][:].tolist())
                counts=np.bincount(codes[codes>=0], minlength=len(cats))
                eligible=[{'gene_id': cat.decode(), 'cells': int(counts[i])} for i,cat in enumerate(cats) if cat in varids and counts[i]>100]
                report['remote_hdf5']={'url': URL, 'size_bytes': remote.size, 'matrix_shape': list(f['X'].shape), 'matrix_dtype': str(f['X'].dtype),
                                      'matrix_offset': f['X'].id.get_offset(), 'contiguous': f['X'].chunks is None,
                                      'control_cells': sum(int(counts[i]) for i,c in enumerate(cats) if c==b'non-targeting'),
                                      'full_cohort_eligible_gene_count_before_strong_effect_filter': len(eligible),
                                      'post_eligibility_directed_pair_count': len(eligible)*(len(eligible)-1),
                                      'eligibility_rule': 'Pinned preprocessing keeps gene_id represented in var with >100 intervention cells; strong-effect filtering not executed',
                                      'full_sha256_verified': False, 'exact_figshare_equivalence': 'UNKNOWN',
                                      'normalization_boundary': 'Exact pinned normalize_per_cell default requires global median full-X row totals; stored UMI_count is not assumed equivalent',
                                      'expression_outcomes_read': False}
            report['range_receipts']=remote.receipts
            report['received_bytes']=remote.received_bytes
        xls=linear_source/'source_data/suppl-qqplot.xlsx'
        if sha(xls)!=XLS_SHA:
            raise ValueError('PUBLISHED_XLS_IDENTITY_MISMATCH')
        wb=openpyxl.load_workbook(xls,read_only=True,data_only=True)
        sheet=wb['Sheet1']
        report['norman_panel']={'source_url': 'https://raw.githubusercontent.com/const-ae/linear_perturbation_prediction-Paper/'+LINEAR+'/source_data/suppl-qqplot.xlsx',
                               'bytes': xls.stat().st_size, 'sha256': XLS_SHA, 'rows_including_header':sheet.max_row, 'columns':sheet.max_column,
                               'column_names': list(next(sheet.iter_rows(min_row=1,max_row=1,values_only=True))),
                               'outcome_values_inspected_by_data_constructor':False,
                               'source_selection': 'Rmd103-109:control-expression rank;391-405:ground_truth A/B/AB/control,seed1,rank<=1000;545:export',
                               'limitations':['1000 source-selected readouts, not full transcriptome','Published condition means, not raw cells or replicate-level uncertainty','No original Figshare equivalence claim','Protected combination split required before builder receives AB or error values']}
        wb.close()
        report['decision']='PIVOT_TO_SOURCE_BOUND_NORMAN_CONDITION_MEANS_PILOT; RPE1_BENCHMARK_UNQUALIFIED'
    except Exception as error:
        report['error']={'type':type(error).__name__,'message':str(error)}
        report['decision']='REFUSED_OR_FAILED_QUALIFICATION'
    report['versions']={p: importlib.metadata.version(p) for p in ('h5py','lightgbm','requests','numpy','scipy','pandas','openpyxl')}
    report['resources']={'wall_seconds':time.monotonic()-started,'self_cpu_seconds':time.process_time()-cpu,
                         'child_cpu_seconds':resource.getrusage(resource.RUSAGE_CHILDREN).ru_utime+resource.getrusage(resource.RUSAGE_CHILDREN).ru_stime,
                         'full_monetary_cost':None,'inference_cost':None}
    (output/'REPORT.json').write_text(json.dumps(report,indent=2)+'\n')
    return report

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',required=True);parser.add_argument('--psgrn-source',type=Path,required=True);parser.add_argument('--linear-source',type=Path,required=True)
    args=parser.parse_args();result=run(args.output,args.psgrn_source,args.linear_source);print(json.dumps({'decision':result['decision'],'resources':result['resources']}))
    raise SystemExit(int('error' in result))
