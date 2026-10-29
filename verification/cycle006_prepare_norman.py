"""Verifier-controlled pair split of the pinned published Norman mean panel."""
import argparse
import hashlib
import json
from pathlib import Path
import secrets
import tempfile
import time

import numpy as np
from openpyxl import load_workbook

SOURCE_SHA = "1b652c352d4dabdd20324d7fd74d331898fa22c51cf6fa28ff65aa1322a6c78f"
SOURCE_COMMIT = "bfa6eeea2bd145a1af2ec0127a2e808cc38456a9"
SOURCE_BLOB = "f35c9efe3bf0f917068b2aaf2c2236b593c06e2c"
SPLIT_SALT = "cycle006-prospective-pair-split-v1"


def sha(raw):
    return hashlib.sha256(raw).hexdigest()


def write(path, value):
    with path.open("x") as stream:
        json.dump(value, stream, indent=2, allow_nan=False)
        stream.write("\n")


def prepare(source, public, private):
    started, cpu = time.monotonic(), time.process_time()
    raw = source.read_bytes()
    if sha(raw) != SOURCE_SHA or hashlib.sha1(b"blob " + str(len(raw)).encode() + b"\0" + raw).hexdigest() != SOURCE_BLOB:
        raise ValueError("PRIMARY_DATA_IDENTITY_MISMATCH")
    public.mkdir(parents=True, exist_ok=False)
    private.mkdir(parents=True, exist_ok=False)
    book = load_workbook(source, read_only=True, data_only=True)
    if len(book.worksheets) != 1:
        raise ValueError("UNEXPECTED_WORKSHEET_SET")
    rows = book.worksheets[0].iter_rows(values_only=True)
    header = tuple(next(rows))
    if header != ("gene_name", "pert_group", "seed", "AB", "A", "B", "ctrl", "error", "rank"):
        raise ValueError("PRIMARY_DATA_SCHEMA_MISMATCH")
    records = {}
    for gene, pair, seed, ab, a, b, control, error, rank in rows:
        if seed != 1 or type(gene) is not str or type(pair) is not str:
            raise ValueError("PRIMARY_PANEL_SCOPE_MISMATCH")
        if gene in records.setdefault(pair, {}):
            raise ValueError("DUPLICATE_PAIR_GENE")
        values = np.asarray([a, b, control, ab], dtype=np.float64)
        if not np.isfinite(values).all():
            raise ValueError("NONFINITE_PRIMARY_DATA")
        records[pair][gene] = values
    book.close()
    if len(records) != 124 or any(len(values) != 1000 for values in records.values()):
        raise ValueError("UNEXPECTED_PAIR_OR_READOUT_COUNT")
    genes = sorted(next(iter(records.values())))
    if any(set(values) != set(genes) for values in records.values()):
        raise ValueError("INCONSISTENT_READOUT_PANEL")
    ordering = sorted(records, key=lambda pair: sha((SOURCE_BLOB + "|" + SPLIT_SALT + "|" + pair).encode()))
    development, heldout = ordering[:80], ordering[80:]
    def arrays(pairs):
        result = np.asarray([[records[pair][gene] for gene in genes] for pair in pairs], dtype=np.float64)
        return {"pairs": np.asarray(pairs), "genes": np.asarray(genes),
                "A": result[:, :, 0], "B": result[:, :, 1], "ctrl": result[:, :, 2], "AB": result[:, :, 3]}
    dev, protected = arrays(development), arrays(heldout)
    dev["fit_indices"] = np.arange(60, dtype=np.int64)
    dev["calibration_indices"] = np.arange(60, 80, dtype=np.int64)
    np.savez_compressed(public / "development.npz", **dev)
    np.savez_compressed(public / "heldout_features.npz", **{key: value for key, value in protected.items() if key != "AB"})
    np.savez_compressed(private / "heldout_truth.npz", **protected)
    public_manifest = {"schema": "agi-alpha-cycle006-norman-protected-split/v1",
        "source_commit": SOURCE_COMMIT, "source_git_blob": SOURCE_BLOB, "source_sha256": SOURCE_SHA,
        "source_bytes": len(raw), "salt": SPLIT_SALT, "split_rule": "Sort pair IDs by SHA256(sourceGitBlob|salt|pair). First80development, next44protected. Within80 first60fit, final20calibration. No AB outcome is used for assignment.",
        "development_pairs": development, "heldout_pairs": heldout, "genes": 1000,
        "public_files": {name: sha((public / name).read_bytes()) for name in ("development.npz", "heldout_features.npz")},
        "protected_truth_sha256": sha((private / "heldout_truth.npz").read_bytes()),
        "disclosed_inputs": ["A", "B", "ctrl", "pair IDs", "alphabetically canonical gene IDs"],
        "withheld_fields": ["heldout AB", "source error", "source rank", "original row ordering"],
        "eligibility": "Published seed1 panel of1000readouts selected by CONTROL expression rank; no arbitrary cell/gene truncation. This retains paper preprocessing and does not reconstruct cell-level variance.",
        "sample_unit": "44 double-intervention pairs;1000shared readoutgenes are not independentreplicates. Sharedsingle inputs and commonCONTROL cause dependence.",
        "protection": "Separate verifier role owns outcome file. Shared filesystem/host does not enforce hostile-user isolation; builder is instructed not to read original workbook or verifier private path. Public paper data may have prior exposure.",
        "resources": {"wall_seconds": time.monotonic() - started, "process_cpu_seconds": time.process_time() - cpu},
        "status": "SPLIT_FROZEN_BEFORE_BUILDER_AB_ACCESS; protected evaluation unrun"}
    write(public / "SPLIT.json", public_manifest)
    write(private / "PRIVATE_LOCATION.json", {"truth": str(private / "heldout_truth.npz"), "source_sha256": SOURCE_SHA,
        "public_split_sha256": sha((public / "SPLIT.json").read_bytes()), "protected_truth_sha256": public_manifest["protected_truth_sha256"]})
    return public_manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--public", type=Path, required=True)
    parser.add_argument("--private", type=Path, required=True)
    args = parser.parse_args()
    result = prepare(args.source, args.public, args.private)
    print(json.dumps({"public": str(args.public), "development_pairs": 80, "heldout_pairs": 44,
        "split_sha256": sha((args.public / "SPLIT.json").read_bytes())}))
